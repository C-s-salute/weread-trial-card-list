# -*- coding: utf-8 -*-
"""把 HTML 同一口径（体验卡可读 + 推荐值 ≥60% + 模糊去重）的书目写成 Excel。

复用 build_html 的过滤与去重函数，保证 Excel 与 HTML 完全一致：
  概览 / 神作榜 / 神作潜力榜 / 好评如潮榜 / 体验卡可读(全部) / 分类汇总

用法：WEREAD_WORK=<数据目录> python build_excel.py
"""
import json
import os
import time

import build_html as m
import wrconfig
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

W = str(wrconfig.WORK)
POOL = '_discover_pool.json'
EXTRA = '_book_extra.json'
READ = '_book_read.json'
OUT = '微信读书_体验卡可读高分书.xlsx'

HEAD_FILL = PatternFill('solid', fgColor='1F3864')
HEAD_FONT = Font(bold=True, color='FFFFFF', size=10, name='微软雅黑')
BODY = Font(size=10, name='微软雅黑')
LINK = Font(size=10, name='微软雅黑', color='0B57B8', underline='single')
CENTER = Alignment(horizontal='center', vertical='center')
LEFT = Alignment(horizontal='left', vertical='center')
THIN = Side(style='thin', color='D9D9D9')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

COLS = ['书名（点击跳转）', '作者', '大类', '子类', '推荐值%', '评价人数',
        '累计阅读人数', '字数', '标价(元)', '出版时间', '出版社', 'bookId']
WIDTHS = [44, 26, 12, 16, 10, 10, 13, 10, 9, 11, 20, 14]

PALETTE = ['FBE2E2', 'FCE8D5', 'FDF3D0', 'EDF7D3', 'DCF0DA', 'D4EDE7',
           'D7EDF5', 'DBE6F6', 'E2E0F5', 'EEE0F4', 'F7DEEB', 'FADDE4',
           'F2E2D0', 'EFEFDE', 'E3EFE3', 'DDEDEF', 'E0E7F2', 'E7E3F1',
           'F1E2EC', 'F5E5DA', 'EBF0D8', 'E2F0EC', 'DEE6EC', 'F0E7E0']


def load():
    payload = json.load(open(os.path.join(W, POOL), encoding='utf-8'))
    extra = json.load(open(os.path.join(W, EXTRA), encoding='utf-8'))
    read = json.load(open(os.path.join(W, READ), encoding='utf-8'))
    raw = [b for b in payload['books'] if b['rights'] == 'T' and b.get('soldout') != 1]
    for b in raw:
        b['_cat'], b['_sub'] = m.split_cat(b)
        b['_read'] = (read.get(b['bookId']) or {}).get('readingCount')
        b['_words'] = (extra.get(b['bookId']) or {}).get('totalWords')
    rated = [b for b in raw if (b.get('rating') or 0) >= 60 and (b.get('ratingCount') or 0) > 0]
    rated.sort(key=lambda b: (-(b.get('rating') or 0), -(b.get('ratingCount') or 0),
                              -(b.get('_read') or 0)))
    for b in rated:
        m.dedup_keys(b)
    kept = []
    for b in rated:
        if any(m.same_book(b, k) for k in kept):
            continue
        kept.append(b)
    return payload, raw, rated, kept


def hdr(ws):
    for i, c in enumerate(COLS, 1):
        cell = ws.cell(row=1, column=i, value=c)
        cell.fill, cell.font, cell.alignment, cell.border = HEAD_FILL, HEAD_FONT, CENTER, BORDER
        ws.column_dimensions[get_column_letter(i)].width = WIDTHS[i - 1]
    ws.row_dimensions[1].height = 28


def fill(ws, books, color_of=None):
    for r, b in enumerate(books, 2):
        vals = [b.get('title') or '', b.get('author') or '', b['_cat'], b['_sub'],
                b.get('rating'), b.get('ratingCount'), b.get('_read'), b.get('_words') or None,
                b.get('price'), b.get('publishTime') or '', b.get('publisher') or '',
                b['bookId']]
        bg = color_of(b) if color_of else None
        for i, v in enumerate(vals, 1):
            cell = ws.cell(row=r, column=i, value=v)
            cell.font, cell.border = BODY, BORDER
            cell.alignment = LEFT if i in (1, 2) else CENTER
            if bg:
                cell.fill = PatternFill('solid', fgColor=bg)
        t = ws.cell(row=r, column=1)
        if b.get('deepLink'):
            t.hyperlink = b['deepLink']
            t.font = LINK
        for i, fmt in ((5, '0.0'), (6, '#,##0'), (7, '#,##0'), (8, '#,##0'), (9, '0.00')):
            ws.cell(row=r, column=i).number_format = fmt
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = 'A1:L%d' % max(len(books) + 1, 2)


def main():
    payload, raw, rated, books = load()
    ts = time.strftime('%Y-%m-%d %H:%M:%S', time.localtime(payload['fetchedAt']))
    n_low, n_dup = len(raw) - len(rated), len(rated) - len(books)

    grp = {}
    for b in books:
        grp.setdefault(b['_cat'], []).append(b)
    cats = sorted(grp, key=lambda k: -len(grp[k]))
    color = {c: PALETTE[i % len(PALETTE)] for i, c in enumerate(cats)}
    color_of = lambda b: color[b['_cat']]

    wb = Workbook()

    # --- 概览 ---
    ws = wb.active
    ws.title = '概览'
    ws['A1'] = '微信读书 · 体验卡可读书 · 高分清单（去重版）'
    ws['A1'].font = Font(bold=True, size=15, name='微软雅黑', color='1F3864')
    ws.merge_cells('A1:C1')
    rows = [
        ('数据抓取时间', ts),
        ('候选池', '%d 本（分类榜/子分类榜 Top 枚举，不等于全库）' % len(payload['books'])),
        ('体验卡可读且未下架', '%d 本' % len(raw)),
        ('剔除无评分 / 推荐值 <60%%', '%d 本' % n_low),
        ('合并同书不同版本', '%d 本' % n_dup),
        ('最终保留', '%d 本 / %d 个大类' % (len(books), len(cats))),
        ('字数已补抓', '%d 本' % len([b for b in books if b.get('_words')])),
    ]
    r = 3
    for k, v in rows:
        ws.cell(row=r, column=1, value=k).font = Font(size=10, name='微软雅黑', color='595959')
        ws.cell(row=r, column=2, value=v).font = Font(bold=True, size=10, name='微软雅黑')
        r += 1

    r += 1
    ws.cell(row=r, column=1, value='三大榜单').font = Font(bold=True, size=11, name='微软雅黑', color='1F3864')
    r += 1
    for name, rule, fn in m.RANKS:
        ws.cell(row=r, column=1, value=name).font = Font(size=10, name='微软雅黑')
        ws.cell(row=r, column=2, value=rule).font = Font(size=10, name='微软雅黑', color='7F7F7F')
        ws.cell(row=r, column=3, value='%d 本' % len([b for b in books if fn(b)])).font = BODY
        r += 1

    r += 1
    for n in [
        '口径说明：',
        '1. 「体验卡可读」= 接口 payingStatus=1（体验卡可免费读付费内容）；已排除 soldout=1 的下架书。',
        '2. 推荐值 = newRating / 10（接口 0~1000 尺度）；评价人数 = ratingCount；'
        '累计阅读人数来自书籍分享页 readingStat.readingCount（列表接口返回的不是累计口径，已改为逐本补抓）。',
        '3. 字数来自书籍详情接口 totalWords，逐本补抓；未抓到的留空。',
        '4. 大类/子类直接取自平台书目元数据 category 字段（形如「历史-历史典籍」），不是按抓取榜单命名。',
        '5. 神作榜 = 推荐值 ≥ 90% 且 累计阅读人数 ≥ 5000；神作潜力榜 = 推荐值 ≥ 90% 且 累计阅读人数 < 5000；'
        '好评如潮榜 = 85% ≤ 推荐值 < 90% 且 累计阅读人数 ≥ 1500。',
        '6. 已剔除无评分（评价人数 0）与推荐值 < 60% 的书。',
        '7. 同一本书的不同版本（书名核心相似 + 作者关键词重合，忽略「（果麦经典）」括注与「[法]」国籍前缀）只保留评分最高的版本；'
        '多卷本（第N卷/（N）/上中下册）与期刊分期（年第N期/增刊/合集）不合并。',
        '8. 「体验卡可读(全部)」按大类（书多在前）→ 子类（本数降序）分块，块内推荐值降序，每个大类一种底色；书名可点击直达微信读书。',
    ]:
        cell = ws.cell(row=r, column=1, value=n)
        cell.font = Font(size=9, name='微软雅黑', bold=(n == '口径说明：'),
                         color='1F3864' if n == '口径说明：' else '7F7F7F')
        r += 1
    ws.column_dimensions['A'].width = 34
    ws.column_dimensions['B'].width = 88
    ws.column_dimensions['C'].width = 12

    # --- 三大榜单 ---
    for name, rule, fn in m.RANKS:
        sel = [b for b in books if fn(b)]
        sel.sort(key=lambda b: (-(b.get('newRating') or 0), -(b.get('_read') or 0)))
        ws_r = wb.create_sheet(name[:31])
        hdr(ws_r)
        fill(ws_r, sel, color_of)

    # --- 全部（大类 → 子类分块）---
    ws3 = wb.create_sheet('体验卡可读(全部)')
    hdr(ws3)
    ordered = []
    for c in cats:
        subs = {}
        for b in grp[c]:
            subs.setdefault(b['_sub'], []).append(b)
        for s in sorted(subs, key=lambda s: (-len(subs[s]), s)):
            subs[s].sort(key=lambda b: (-(b.get('newRating') or 0), -(b.get('_read') or 0)))
            ordered.extend(subs[s])
    fill(ws3, ordered, color_of)

    # --- 分类汇总 ---
    ws4 = wb.create_sheet('分类汇总')
    cols4 = ['大类', '本数', '子类数', '推荐值≥90%', '85%≤推荐值<90%', '推荐值<70%', '有阅读人数']
    for i, c in enumerate(cols4, 1):
        cell = ws4.cell(row=1, column=i, value=c)
        cell.fill, cell.font, cell.alignment, cell.border = HEAD_FILL, HEAD_FONT, CENTER, BORDER
    for i, wd in enumerate([16, 9, 9, 13, 18, 13, 12], 1):
        ws4.column_dimensions[get_column_letter(i)].width = wd
    for rr, c in enumerate(cats, 2):
        rs = grp[c]
        vals = [c, len(rs), len({b['_sub'] for b in rs}),
                len([b for b in rs if (b.get('rating') or 0) >= 90]),
                len([b for b in rs if 85 <= (b.get('rating') or 0) < 90]),
                len([b for b in rs if (b.get('rating') or 0) < 70]),
                len([b for b in rs if b.get('_read')])]
        for i, v in enumerate(vals, 1):
            cell = ws4.cell(row=rr, column=i, value=v)
            cell.font, cell.border = BODY, BORDER
            cell.alignment = LEFT if i == 1 else CENTER
        ws4.cell(row=rr, column=1).fill = PatternFill('solid', fgColor=color[c])
    ws4.freeze_panes = 'A2'

    wb.save(os.path.join(W, OUT))
    print('已写出 %s（%d 本 / %d 类 / 榜单 %d）'
          % (os.path.abspath(os.path.join(W, OUT)), len(books), len(cats), len(m.RANKS)), flush=True)


if __name__ == '__main__':
    main()
